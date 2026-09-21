#!/usr/bin/env python3
"""课堂分享 PPT：架构 → 库/ER → 模块 → 部署运维。"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT = "/Users/andrew_file/Downloads/github_demo/warehouse/docs/course/10-课堂分享.pptx"

BG = RGBColor(0x0B, 0x0D, 0x12)
CARD = RGBColor(0x16, 0x1A, 0x22)
LINE = RGBColor(0x2A, 0x31, 0x3D)
GOLD = RGBColor(0xC9, 0xA2, 0x27)
GOLD2 = RGBColor(0xE8, 0xD5, 0x8A)
RED = RGBColor(0xC4, 0x1E, 0x3A)
IVORY = RGBColor(0xF5, 0xF2, 0xEA)
MUTED = RGBColor(0x9A, 0xA0, 0xAB)

FONT = "PingFang SC"
W, H = Inches(13.333), Inches(7.5)
TOTAL = 14


def set_run(run, size=18, color=IVORY, bold=False, font=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(r_pr, qn("a:ea"))
    ea.set("typeface", font)


def fill_text(shape, lines, size=16, color=IVORY, bold=False, align=PP_ALIGN.LEFT, space=8):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.clear()
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space)
        run = p.add_run()
        run.text = line
        set_run(run, size, color, bold)
    return shape


def solid(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def stroke(shape, color: RGBColor, pt=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(pt)


def rect(slide, l, t, w, h, fill=CARD, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    solid(s, fill)
    if line:
        stroke(s, line, 1.15)
    else:
        s.line.fill.background()
    try:
        s.adjustments[0] = 0.08
    except Exception:
        pass
    return s


def bar(slide, l, t, w, h, fill=GOLD):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    solid(s, fill)
    return s


def set_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def header(slide, kicker: str, title: str, sub: str | None = None):
    bar(slide, Inches(0.7), Inches(0.38), Inches(0.08), Inches(0.42), GOLD)
    k = slide.shapes.add_textbox(Inches(0.95), Inches(0.28), Inches(11.5), Inches(0.28))
    fill_text(k, [kicker], 11, GOLD, True)
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.52), Inches(12), Inches(0.55))
    fill_text(t, [title], 28, IVORY, True)
    if sub:
        s = slide.shapes.add_textbox(Inches(0.7), Inches(1.05), Inches(12), Inches(0.35))
        fill_text(s, [sub], 13, MUTED, False)


def footer(slide, page: int):
    bar(slide, Inches(0.7), Inches(7.12), Inches(11.9), Pt(0.75), LINE)
    f = slide.shapes.add_textbox(Inches(0.7), Inches(7.16), Inches(9), Inches(0.25))
    fill_text(f, ["药企进销存  ·  课堂分享"], 10, MUTED)
    n = slide.shapes.add_textbox(Inches(10.4), Inches(7.16), Inches(2.2), Inches(0.25))
    fill_text(n, [f"{page:02d}  /  {TOTAL:02d}"], 10, GOLD, True, PP_ALIGN.RIGHT)


def card(slide, l, t, w, h, title, body_lines, accent=GOLD):
    box = rect(slide, l, t, w, h, CARD, LINE)
    bar(slide, l, t, Inches(0.07), h, accent)
    tit = slide.shapes.add_textbox(l + Inches(0.22), t + Inches(0.12), w - Inches(0.35), Inches(0.36))
    fill_text(tit, [title], 15, GOLD2, True)
    body = slide.shapes.add_textbox(l + Inches(0.22), t + Inches(0.48), w - Inches(0.35), h - Inches(0.58))
    fill_text(body, body_lines, 13, IVORY, False, space=6)
    return box


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    def ns():
        s = prs.slides.add_slide(blank)
        set_bg(s)
        return s

    # 1 cover
    s = ns()
    bar(s, Inches(0), Inches(0), Inches(13.333), Inches(0.08), RED)
    bar(s, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08), GOLD)
    k = s.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11), Inches(0.4))
    fill_text(k, ["CLASS  ·  ARCHITECTURE  ·  DATABASE  ·  MODULES  ·  OPS"], 13, GOLD, True)
    t = s.shapes.add_textbox(Inches(1.1), Inches(2.3), Inches(11), Inches(1.5))
    fill_text(t, ["药企进销存", "课堂分享"], 44, IVORY, True, space=6)
    sub = s.shapes.add_textbox(Inches(1.1), Inches(4.6), Inches(10.5), Inches(1.0))
    fill_text(sub, ["架构 → 库 / ER → 模块 → 部署运维", "B/S  ·  批号库存  ·  追溯  ·  日清月结"], 16, MUTED, False, space=8)

    # 2 agenda
    s = ns()
    header(s, "00  /  目录", "今天按老师提纲讲四段", "不报功能清单，把结构和账讲清楚")
    items = [
        ("01", "架构", "谁访问谁，请求怎么落到库"),
        ("02", "库 / ER", "批号粒度，表怎么切"),
        ("03", "模块", "进、存、出、码、结账"),
        ("04", "部署运维", "怎么上线、备份、CI"),
    ]
    for i, (no, title, desc) in enumerate(items):
        x = Inches(0.7) + i * Inches(3.15)
        card(s, x, Inches(1.7), Inches(3.0), Inches(4.8), f"{no}  {title}", [desc, "", "各约 3 分钟。"], GOLD if i % 2 == 0 else RED)
    footer(s, 2)

    # 3 position
    s = ns()
    header(s, "01  /  架构", "先定系统是干什么的", "药品流通进销存，不是门店收银台")
    card(s, Inches(0.7), Inches(1.65), Inches(6.0), Inches(5.1), "要回答的三句话", [
        "货在哪：某仓、某批、合格还是待验。",
        "票开了没：出库单带发票号。",
        "钱回来没：未回 / 部分 / 已回。",
        "",
        "再加一句：码能否按发票或 SPDID 查到。",
    ], RED)
    card(s, Inches(6.95), Inches(1.65), Inches(5.65), Inches(5.1), "本期主链路", [
        "品种 / 供应商 / 客户",
        "采购入库，确认后入账",
        "批号库存",
        "追溯码",
        "销售出库，不足拒绝",
        "医院收货",
        "日清 / 月结",
    ])
    footer(s, 3)

    # 4 logical arch
    s = ns()
    header(s, "01  /  架构", "逻辑分层", "客户端不碰库存 SQL")
    layers = [
        ("客户端", "Vue 3 Web、uni-app X、Flutter、Electron"),
        ("接入", "Nginx：静态页走 / ，接口走 /warehouse"),
        ("应用", "Controller → Service 事务 → Mapper"),
        ("横切", "SA-Token 权限、操作日志 AOP"),
        ("数据", "MySQL pharma_ims；sys_* 与药企表同库分开"),
    ]
    for i, (title, desc) in enumerate(layers):
        y = Inches(1.58) + i * Inches(1.02)
        rect(s, Inches(0.7), y, Inches(12.0), Inches(0.92), CARD, LINE)
        bar(s, Inches(0.7), y, Inches(0.08), Inches(0.92), GOLD if i != 3 else RED)
        n = s.shapes.add_textbox(Inches(1.05), y + Inches(0.12), Inches(2.4), Inches(0.7))
        fill_text(n, [title], 18, GOLD2, True)
        d = s.shapes.add_textbox(Inches(3.6), y + Inches(0.22), Inches(8.8), Inches(0.55))
        fill_text(d, [desc], 15, IVORY)
    footer(s, 4)

    # 5 deploy arch
    s = ns()
    header(s, "01  /  架构", "部署就三块", "浏览器只认识 80/443")
    card(s, Inches(0.7), Inches(1.65), Inches(3.9), Inches(5.1), "Nginx", [
        "location /  → 前端 dist",
        "location /warehouse → 127.0.0.1:8899",
        "",
        "生产再加 443 与跳转。",
    ], RED)
    card(s, Inches(4.8), Inches(1.65), Inches(3.9), Inches(5.1), "Spring Boot", [
        "JAR 听 8899",
        "上下文 /warehouse",
        "确认入库/出库开事务",
        "",
        "不把口令写进仓库。",
    ])
    card(s, Inches(8.9), Inches(1.65), Inches(3.7), Inches(5.1), "MySQL", [
        "库名 pharma_ims",
        "档案、单据、库存、码、日清月结",
        "",
        "备份用 deploy/backup.sh",
    ])
    footer(s, 5)

    # 6 db principle
    s = ns()
    header(s, "02  /  库", "账不是一个商品一个数", "药品 × 仓库 × 批号 × 质量状态")
    box = rect(s, Inches(0.7), Inches(1.7), Inches(12.0), Inches(1.6), CARD, LINE)
    q = s.shapes.add_textbox(Inches(1.0), Inches(2.05), Inches(11.4), Inches(1.0))
    fill_text(q, ["uk_batch_stocks (drug_id, warehouse_id, batch_no, quality_status)"], 22, GOLD2, True, PP_ALIGN.CENTER)
    card(s, Inches(0.7), Inches(3.55), Inches(3.9), Inches(3.15), "进货确认", ["按入库数量 increase", "草稿不动库存"], GOLD)
    card(s, Inches(4.8), Inches(3.55), Inches(3.9), Inches(3.15), "出库确认", ["按批号 decrease", "不够就拒绝"], RED)
    card(s, Inches(8.9), Inches(3.55), Inches(3.7), Inches(3.15), "红冲 / 升益", ["红冲另开新单再加回", "升益按实盘减账面"], GOLD)
    footer(s, 6)

    # 7 ER layers
    s = ns()
    header(s, "02  /  ER", "表按六层切，不要画成一张蜘蛛网", "网页 /design 有完整 ER")
    layers = [
        ("档案", "药品、客户、供应商、仓库、员工"),
        ("进货", "进货单 + 明细，含批号效期"),
        ("出库", "出库单 + 明细，发票号、红冲、回款、收货"),
        ("库存", "批号库存、升益单"),
        ("追溯", "追溯码、包装关联，靠 SPDID 对齐"),
        ("结账", "日清、月结快照，不挂单据外键"),
    ]
    for i, (title, desc) in enumerate(layers):
        col, row = i % 3, i // 3
        x = Inches(0.7) + col * Inches(4.15)
        y = Inches(1.65) + row * Inches(2.55)
        card(s, x, y, Inches(3.95), Inches(2.35), title, [desc])
    footer(s, 7)

    # 8 key relations
    s = ns()
    header(s, "02  /  ER", "几条关系就够答辩", "完整图在 docs/pharma-er-uml.md")
    rows = [
        ("供应商 1—N 进货单", "一单一个上游"),
        ("客户 1—N 出库单", "一单一个医院或药店"),
        ("进货/出库 1—N 明细", "至少一行才能确认"),
        ("明细 1—N 追溯码", "逻辑外键 SPDID"),
        ("红冲 N—1 原出库", "original_invoice_no，不改原单"),
        ("旧仓管 bus_*", "按 SKU 扣库存，禁止混进药企图"),
    ]
    for i, (a, b) in enumerate(rows):
        y = Inches(1.58) + i * Inches(0.85)
        rect(s, Inches(0.7), y, Inches(12.0), Inches(0.75), CARD, LINE)
        left = s.shapes.add_textbox(Inches(1.0), y + Inches(0.16), Inches(5.5), Inches(0.45))
        fill_text(left, [a], 16, GOLD2, True)
        right = s.shapes.add_textbox(Inches(6.6), y + Inches(0.16), Inches(5.8), Inches(0.45))
        fill_text(right, [b], 15, IVORY)
    footer(s, 8)

    # 9 modules
    s = ns()
    header(s, "03  /  模块", "后端仍是 Controller → Service → Mapper", "sys 管人，bus 管货")
    mods = [
        ("系统", "用户、角色、部门、菜单、日志"),
        ("档案", "品种、客户、供应商、仓库"),
        ("进货", "草稿、确认入库、发票查询"),
        ("出库", "开票出库、红冲、回款、收货"),
        ("仓储", "批号库存、盘点、调拨、升益"),
        ("追溯", "查码、扫码纠错、大码解析"),
        ("结账", "日清、月结"),
        ("报表", "进销金额/商品/利润、流向"),
    ]
    for i, (title, desc) in enumerate(mods):
        col, row = i % 4, i // 4
        x = Inches(0.7) + col * Inches(3.15)
        y = Inches(1.65) + row * Inches(2.55)
        card(s, x, y, Inches(3.0), Inches(2.35), title, [desc], RED if title in ("出库", "追溯") else GOLD)
    footer(s, 9)

    # 10 main flow
    s = ns()
    header(s, "03  /  模块", "主链路压成一句", "草稿不入账，确认才动库存")
    steps = [
        ("进", "采购开草稿"),
        ("入", "仓管确认"),
        ("存", "批号库存"),
        ("出", "销售开票"),
        ("扣", "确认出库"),
        ("收", "医院签收"),
        ("清", "日清月结"),
    ]
    for i, (name, desc) in enumerate(steps):
        x = Inches(0.55) + i * Inches(1.82)
        rect(s, x, Inches(2.0), Inches(1.68), Inches(2.15), CARD, LINE)
        n = s.shapes.add_textbox(x, Inches(2.2), Inches(1.68), Inches(0.7))
        fill_text(n, [name], 28, GOLD2, True, PP_ALIGN.CENTER)
        d = s.shapes.add_textbox(x, Inches(3.0), Inches(1.68), Inches(0.9))
        fill_text(d, [desc], 13, IVORY, False, PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            ar = s.shapes.add_textbox(x + Inches(1.5), Inches(2.7), Inches(0.4), Inches(0.5))
            fill_text(ar, ["→"], 18, GOLD, True, PP_ALIGN.CENTER)
    note = s.shapes.add_textbox(Inches(0.7), Inches(4.5), Inches(12.0), Inches(2.1))
    fill_text(note, [
        "库存不足：出库确认失败，草稿可删，账是干净的。",
        "收货重复提交：拒绝。红冲是新单，不是改蓝字。",
        "追溯：发票号看整单，SPDID 或扫码看单品。",
    ], 16, MUTED, False, space=10)
    footer(s, 10)

    # 11 deploy steps
    s = ns()
    header(s, "04  /  部署", "上线按这个顺序", "先备份，再换包，最后巡检")
    steps = [
        ("1", "备份", "deploy/backup.sh 压一份 sql.gz"),
        ("2", "后端", "mvn 打 JAR，重启 8899"),
        ("3", "前端", "npm run build，覆盖 Nginx dist"),
        ("4", "冒烟", "登录 → 批号库存 → 一张草稿出入库"),
        ("5", "失败则回滚", "换回上一版 JAR / dist，或 restore.sh"),
    ]
    for i, (no, title, desc) in enumerate(steps):
        y = Inches(1.58) + i * Inches(1.0)
        rect(s, Inches(0.7), y, Inches(12.0), Inches(0.88), CARD, LINE)
        num = s.shapes.add_textbox(Inches(0.95), y + Inches(0.2), Inches(0.6), Inches(0.5))
        fill_text(num, [no], 20, GOLD, True)
        tt = s.shapes.add_textbox(Inches(1.7), y + Inches(0.2), Inches(2.4), Inches(0.5))
        fill_text(tt, [title], 18, IVORY, True)
        dd = s.shapes.add_textbox(Inches(4.3), y + Inches(0.22), Inches(8.0), Inches(0.5))
        fill_text(dd, [desc], 15, MUTED)
    footer(s, 11)

    # 12 ops
    s = ns()
    header(s, "04  /  运维", "日常三件事", "脚本在 deploy/，口令只走环境变量")
    card(s, Inches(0.7), Inches(1.65), Inches(3.9), Inches(5.1), "备份", [
        "每天 mysqldump 压缩",
        "保留最近 14 份",
        "不放在 Web 根目录",
        "",
        "backup.sh / restore.sh",
    ], RED)
    card(s, Inches(4.8), Inches(1.65), Inches(3.9), Inches(5.1), "巡检", [
        "首页 200",
        "公钥接口 code=200",
        "Java / Nginx 在跑",
        "磁盘 < 85%，备份未过期",
        "",
        "healthcheck.sh",
    ])
    card(s, Inches(8.9), Inches(1.65), Inches(3.7), Inches(5.1), "生产门槛", [
        "备案域名 + HTTPS",
        "岗位账号，废止弱口令",
        "IP + admin/123456",
        "不能当交付。",
    ])
    footer(s, 12)

    # 13 CI
    s = ns()
    header(s, "04  /  CI", "持续集成：推代码就自动验", "不是“网站已经打开”")
    card(s, Inches(0.7), Inches(1.65), Inches(6.0), Inches(5.1), "这条流水线做什么", [
        "文件：.github/workflows/ci.yml",
        "触发：push、Pull Request",
        "后端：mvn -DskipTests=false test",
        "前端：npm ci && npm run build",
        "",
        "通过只说明能编译、测试、打包。",
        "不自动发到 123.57.134.57。",
    ], RED)
    card(s, Inches(6.95), Inches(1.65), Inches(5.65), Inches(5.1), "和部署的关系", [
        "CI = 持续集成，发现有人把主分支弄坏。",
        "CD = 再往后自动发布。",
        "",
        "我们现在：CI 有，发布仍是手工。",
        "短文：docs/course/技术选型.md",
        "      docs/course/架构设计.md",
    ])
    footer(s, 13)

    # 14 end
    s = ns()
    bar(s, Inches(0), Inches(0), Inches(13.333), Inches(0.08), RED)
    bar(s, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08), GOLD)
    k = s.shapes.add_textbox(Inches(1.1), Inches(2.1), Inches(11), Inches(0.4))
    fill_text(k, ["要看图，打开网页"], 14, GOLD, True)
    t = s.shapes.add_textbox(Inches(1.1), Inches(2.55), Inches(11), Inches(1.4))
    fill_text(t, ["localhost:8888/design", "未登录也能看齐 ER / 用例 / 时序"], 32, IVORY, True, space=10)
    sub = s.shapes.add_textbox(Inches(1.1), Inches(4.7), Inches(10.5), Inches(1.2))
    fill_text(sub, [
        "短文：docs/course/技术选型.md、架构设计.md",
        "运维：deploy/ 与 docs/delivery/05-运维交付包.md",
    ], 16, MUTED, False, space=8)

    prs.save(OUT)
    print("saved", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    try:
        build()
    except ModuleNotFoundError:
        from pptx_stdlib import build as fallback
        path = fallback()
        print("python-pptx 不可用，已用标准库写出", path)
